"""Build a Traditional Chinese character learnset snapshot from the public sheet.

The spreadsheet is community-maintained.  Only rank columns are imported; chat,
effect notes, and the summary tables beneath the character rows are deliberately
excluded.  English source names are retained beside Chinese working translations.
"""

from __future__ import annotations

import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import urlopen

import openpyxl


ROOT = Path(__file__).resolve().parent
SHEET_ID = "1YW5AdvPUbLPr1RAGlnotcRaNTFCQPTKiIgshwrcdUtE"
SOURCE_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/htmlview#gid=1696867101"
EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"

PROFICIENCIES = {
    "Sword": "劍術",
    "Spear": "槍術",
    "Axe": "斧術",
    "Bow": "弓術",
    "Gauntlet": "格鬥術",
    "Faith": "白魔法",
    "Reason": "黑魔法",
    "Authority": "指揮術",
    "Infantry": "步兵術",
    "Riding": "騎術",
    "Armor": "重裝術",
    "Flying": "飛行術",
}

# The sheet and the field guide use different romanisations for eight characters.
CHARACTER_ALIASES = {
    "Centurio": "Centurion",
    "Creek": "Kiryk",
    "Eshmel": "Savior",
    "Hong Hua": "Kouka",
    "Orchel": "Orhel",
    "Tahonia": "Tahounia",
    "Talimun": "Talimoon",
    "Troy": "Troia",
}

NAME_FIXES = {
    "MIghty Fist": "Mighty Fist",
    "Sprit Blow": "Spirit Blow",
    "Defensive Formation+": "Defense Formation+",
    "Lend Defense +": "Lend Defense+",
    "Lend Strength +": "Lend Strength+",
    "Lend Speed +": "Lend Speed+",
    "Lend Speed +?": "Lend Speed+",
    "Lend Dex+": "Lend Dexterity+",
    "Rein +": "Rein+",
    "Threaten +": "Threaten+",
    "Close Quaters": "Close Quarters",
    "Close quarters+": "Close Quarters+",
    "Steady Shot +": "Steady Shot+",
    "Armored Move +": "Armored Move+",
    "Bowblock +": "Bowblock+",
    "Build +5": "Build+5",
    "Build +10": "Build+10",
    "Counterfurry": "Counterflurry",
    "Narrow Defence": "Narrow Defense",
    "Fitting Flier": "Flitting Flier",
    "Fitting Flier+": "Flitting Flier+",
    "Flitting Flier?": "Flitting Flier",
    "Airfighter?": "Airfighter",
    "Cavalry Move?": "Cavalry Move",
    "Armored Move*": "Armored Move",
    "Bowblock+*": "Bowblock+",
    "Beat Down": "Beatdown",
    "Mind's-Eye Blow": "Mind's Eye Blow",
}

TRANSLATIONS = {
    # Sword
    "Blade Slam": "劍身猛擊", "Blaze Blow": "烈焰斬", "Chop": "劈斬",
    "Crosswise Cut": "十字斬", "Dark Ray": "暗黑光線", "Destreza": "精妙劍技",
    "Fateful Strike": "宿命一擊", "Finesse Blade": "技巧之刃", "Flash Blow": "閃光斬",
    "Furious Strike": "怒濤斬", "Gale Blow": "疾風斬", "Greater Hexblade": "高階魔刃",
    "Grounder": "擊墜斬", "Haze Slice": "霧霞斬", "Heat Haze": "陽炎",
    "Hexblade": "魔刃", "Lunging Gale": "突進疾風", "Mist Reaper": "霧之收割",
    "Rend Armor": "破甲斬", "Ripple Slash": "波紋斬", "Shadow Gambit": "影之奇策",
    "Single Sunder": "一刀破碎", "Skysplit": "裂空斬", "Subdue": "制伏",
    "Sunder": "破碎斬", "Swordfaire": "劍術專精", "Tiger Slash": "虎斬",
    "Twin Blades": "雙刃斬", "Vicious Flash": "惡閃", "Windsweep": "風掃",
    "Wrath Strike": "憤怒一擊",
    # Spear
    "Armored Blow": "鎧甲一擊", "Brilliant Strike": "輝煌一擊", "Defending Spear": "守護槍",
    "Double Throw": "雙重投擲", "Flowing Spear": "流水槍", "Flying Spear": "飛槍",
    "Focused Throw": "集中投擲", "Frozen Spear": "冰槍", "Frozen-Spear Slice": "冰槍斬",
    "Gale Spear": "疾風槍", "Graceful Blow": "優雅一擊", "Holy Spear": "聖槍",
    "Jab": "刺擊", "Khamsin": "卡姆辛", "Knightkneel": "屈騎一擊",
    "Brilliant Blow": "輝煌猛擊",
    "Protective Spear": "防護槍", "Protective Storm": "防護風暴", "Quick Thrust": "快速突刺",
    "Raging Tempest": "狂暴風暴", "Spearfaire": "槍術專精", "Stinger": "毒刺",
    "Sword Breaker": "破劍者", "Tempest Spear": "風暴槍", "Vanguard Spear": "先鋒槍",
    "Yawning Throw": "破綻投擲",
    # Axe
    "Battleblade Strike": "戰刃一擊", "Bloodbath": "血浴", "Careful Attack": "慎重攻擊",
    "Crash": "崩擊", "Devilwind Flurry": "魔風亂舞", "Devilwind Fury": "魔風之怒",
    "Discerning Blow": "洞察一擊", "Double Attack": "雙重攻擊", "Double Swing": "雙重揮擊",
    "Enrage": "激怒", "Focused Strike": "集中一擊", "Giant Breaker": "巨人破壞者",
    "Guard Spike": "防守尖擊", "Half Spike": "半尖擊", "Helm Splitter": "破盔",
    "Hurl": "投擲", "Instant Strike": "瞬擊", "Killer Spike": "必殺尖擊",
    "Lightning Axe": "雷斧", "Magic Breaker": "破魔者", "Magic Throw": "魔力投擲",
    "Power Spike": "強力尖擊", "Power Throw": "強力投擲", "Reflecting Dance": "反射之舞",
    "Riot": "暴動", "Smash": "猛擊", "Spear Buster": "破槍者", "Spike": "尖擊",
    "Triple Spike": "三連尖擊", "Wild Abandon": "捨身猛攻",
    # Bow
    "Arrowcrack": "破箭", "Breakaway Shot": "脫離射擊", "Calm Call": "沉著呼喚",
    "Curved Shot": "曲射", "Encloser": "封鎖射擊", "Haphazard": "亂射",
    "Lead Arrow": "先導箭", "Leaden Draw": "沉重拉弓", "Leg Shot": "腿部射擊",
    "Magic Arrow": "魔法箭", "Mighty Shot": "強力射擊", "Shadow Arrow": "影箭",
    "Sneaky Shot": "偷襲射擊", "Snipe": "狙擊", "Spiral Shot": "螺旋射擊",
    "Spirit Bow": "靈力弓", "Surprise Shot": "奇襲射擊", "Twin-Shot Flow": "雙射流",
    "Vital Shot": "要害射擊", "Warning Shot": "警告射擊",
    # Gauntlet
    "Assassin's Fist": "暗殺拳", "Battle Blow": "戰鬥一擊", "Battle Rush": "戰鬥突進",
    "Beatdown": "痛擊", "Cautious Blow": "謹慎一擊", "Crush Armor": "破甲拳",
    "Curtail Magic": "抑制魔力", "Deal Death": "致命一擊", "Fading Assault": "消隱猛攻",
    "Fading Blow": "消隱一擊", "Feint": "佯攻", "Jack Hit": "傑克一擊",
    "Knockout": "擊倒", "Mighty Fist": "強力拳", "Mind's Eye Blow": "心眼一擊",
    "Mystic Blow": "奧秘一擊", "One-Two Punch": "一二連拳", "One-on-One": "單挑",
    "Power Combo": "強力連擊", "Razor Fist": "剃刀拳", "Revival": "復甦",
    "Saving Fist": "救命拳", "Spirit Blow": "靈力一擊", "Spirit Fist": "靈力拳",
    # Faith
    "Aura": "光環", "Barrier": "屏障", "Berserk": "狂暴", "Emergis": "埃梅爾吉斯",
    "Fortify": "全體聖療", "Freeze": "冰封", "Heal": "治療", "Laia": "萊亞",
    "Nosferatu": "吸星術", "Physic": "遠程聖療", "Recover": "高階聖療",
    "Regenera": "再生", "Rescue": "救援", "Resquia": "小型救援", "Restore": "狀態恢復",
    "Seraphim": "熾天使", "Silence": "沉默", "Torch": "火炬", "Ward": "魔防屏障",
    "Warp": "傳送",
    # Reason
    "Blizzard": "暴風雪", "Bolganone": "火山爆焰", "Bolting": "遠雷",
    "Cutting Gale": "裂空風刃", "Dark Spikes T": "暗黑尖刺T", "Death": "死亡",
    "Excalibur": "聖劍風暴", "Fire": "火焰", "Glass Wheel": "玻璃之輪",
    "Ice Blade": "冰刃", "Luna": "月神", "Miasma": "瘴氣", "Mire": "沼澤",
    "Sagittae": "光箭", "Thoron": "雷神", "Thunder": "雷電", "Wandering Wall": "徘徊之壁",
    "Wind": "風刃",
    # Authority
    "Attack Formation": "攻擊陣形", "Barging Order": "突進號令", "Counter Cavalry": "反騎兵",
    "Counter Mages": "反魔法", "Counter Order": "反擊號令", "Counter Snipers": "反狙擊",
    "Defense Formation": "防禦陣形", "Distraction": "聲東擊西", "Draw Back": "拉回",
    "Frontline Gambit": "前線計策", "Gambit Mastery": "計策精通", "Healing Order": "治療號令",
    "Lend Defense": "援防", "Lend Dexterity": "援技", "Lend Luck": "援運",
    "Lend Magic": "援魔", "Lend Resistance": "援魔防", "Lend Speed": "援速",
    "Lend Strength": "援力", "Power-Arts Order": "戰技強化號令", "Propel": "推進",
    "Rein": "牽制", "Safety First Order": "安全第一號令", "Shove": "推擊", "Swap": "交換",
    "Teach Monkly Havoc": "傳授僧侶之亂", "Teach One Chance": "傳授一線生機",
    "Threaten": "威嚇",
    # Infantry
    "Avoid Fatality": "避免致命", "Close Quarters": "近身戰", "Close-Call Evasion": "險境迴避",
    "Combat Advantage": "戰鬥優勢", "Dodge": "閃身", "Evasion": "迴避",
    "Infantry Canto": "步兵再移動", "Infantry Move": "步兵移動", "Pass": "穿越",
    "Positioning": "站位", "Sniper Evasion": "狙擊迴避", "Steady Shot": "穩定射擊",
    "Surprise Attack": "奇襲",
    # Riding
    "Cavalry Move": "騎兵移動", "Knightly Senses": "騎士直覺", "Vanguard": "先鋒",
    # Armor
    "Absolute Defense": "絕對防禦", "Armored Combat": "重裝戰鬥", "Armored Leap": "重裝跳躍",
    "Armored Move": "重裝移動", "Bowblock": "弓箭格擋", "Build": "體格",
    "Counterflurry": "反擊連打", "Crit Defense": "必殺防禦", "Defense Fend": "防禦抵抗",
    "Defense Stance": "防禦架勢", "Dexterity Fend": "技巧抵抗", "Healing Armor": "治療裝甲",
    "Heavy Break": "重擊破壞", "Luck Fend": "幸運抵抗", "Magic Fend": "魔力抵抗",
    "Narrow Defense": "窄域防禦", "Pursuit Defense": "追擊防禦", "Resist Stance": "魔防架勢",
    "Resistant Armor": "抗魔裝甲", "Spearblock": "槍擊格擋", "Strength Fend": "力量抵抗",
    # Flying
    "Airfighter": "空戰者", "Flitting Flier": "靈巧飛騎", "Read the Winds": "讀風",
}

SKIP_VALUES = {"", "?", "-", "---", "N/A", "None"}
RANK_PATTERN = re.compile(r"^[A-S](?:\+)?$")


def split_cell(value: object, default_rank: str):
    raw = str(value or "").strip()
    if not raw or raw in SKIP_VALUES:
        return []
    uncertain = "?" in raw
    # Dark-magic entries append long comma-separated scroll-location notes.
    # Keep the spell name and discard the note before ordinary item splitting.
    if "(Scroll" in raw:
        raw = raw.split("(Scroll", 1)[0].strip()
    # Slashes are item separators only in these known compound cells.
    raw = raw.replace("Lend Defense+/Lend Strength ++", "Lend Defense+, Lend Strength++")
    raw = re.sub(r"\b(Fire|Thunder|Wind)/Wandering Wall\b", r"\1, Wandering Wall", raw)
    pieces = re.split(r"[,\n]+", raw)
    results = []
    for piece in pieces:
        text = piece.strip()
        if not text or text in SKIP_VALUES or text.lower().startswith("scroll:"):
            continue
        # Dark spell cells append acquisition notes after the spell name.
        text = re.split(r"\s*\(Scroll\b", text, maxsplit=1)[0].strip()
        text = re.sub(r"\s+[Α-Ω]$", "", text.replace("�", "")).strip()
        rank = default_rank
        level_match = re.search(r"\(Lv\.\s*(\d+)\)\s*$", text, re.I)
        if level_match:
            rank = f"角色 Lv.{level_match.group(1)}"
            text = text[:level_match.start()].strip()
        else:
            override = re.search(r"\((base|[A-S]\+?)\)\s*$", text, re.I)
            if override:
                token = override.group(1)
                rank = "基礎" if token.lower() == "base" else token.upper()
                text = text[:override.start()].strip()
        # One source cell embeds an additional S+ skill in parentheses.
        extra = None
        extra_match = re.search(r"\((Airfighter\+)\s+(S\+)\)\s*$", text, re.I)
        if extra_match:
            extra = (extra_match.group(1), extra_match.group(2))
            text = text[:extra_match.start()].strip()
        text = NAME_FIXES.get(text, text).strip().rstrip("?").strip()
        if text and text not in SKIP_VALUES:
            results.append((text, rank, uncertain))
        if extra:
            results.append((extra[0], extra[1], uncertain))
    return results


def translate(name: str) -> str:
    suffix = ""
    base = name
    suffix_match = re.search(r"(\+{1,2}|\+\d+)$", name)
    if suffix_match:
        suffix = suffix_match.group(1)
        base = name[:suffix_match.start()].strip()
    translated = TRANSLATIONS.get(base)
    return f"{translated}{suffix}" if translated else ""


def main() -> None:
    characters = json.loads((ROOT / "characters.json").read_text(encoding="utf-8"))
    by_name = {item["name"]: item for item in characters}
    source_to_canonical = {name: name for name in by_name}
    source_to_canonical.update(CHARACTER_ALIASES)

    with urlopen(EXPORT_URL, timeout=60) as response:
        workbook_bytes = response.read()
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), data_only=True, read_only=True)

    entries = []
    seen = set()
    sheet_counts = {}
    for sheet_name, proficiency_zh in PROFICIENCIES.items():
        sheet = workbook[sheet_name]
        rank_columns = []
        for column in range(2, 10):
            header = str(sheet.cell(1, column).value or "").strip()
            if RANK_PATTERN.fullmatch(header):
                rank_columns.append((column, header))
        matched_characters = set()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            source_character = str((row[0] if row else None) or "").strip()
            character = source_to_canonical.get(source_character)
            if not character:
                continue
            matched_characters.add(character)
            metadata = by_name.get(character, {})
            character_zh = (metadata.get("personalInfo") or {}).get("zh") or character
            for column, default_rank in rank_columns:
                value = row[column - 1] if len(row) >= column else None
                for english, rank, uncertain in split_cell(value, default_rank):
                    key = (character, proficiency_zh, rank, english)
                    if key in seen:
                        continue
                    seen.add(key)
                    entries.append({
                        "character": character,
                        "character_zh": character_zh,
                        "source_character": source_character,
                        "proficiency": proficiency_zh,
                        "source_sheet": sheet_name,
                        "rank": rank,
                        "skill_zh": translate(english),
                        "skill_en": english,
                        "uncertain": uncertain,
                        "source_cell": str(value).strip(),
                    })
        sheet_counts[sheet_name] = len(matched_characters)

    rank_order = {rank: index for index, rank in enumerate(
        ["基礎", "D", "D+", "C", "C+", "B", "B+", "A", "A+", "S", "S+", "角色 Lv.45", "角色 Lv.48"]
    )}
    proficiency_order = {name: index for index, name in enumerate(PROFICIENCIES.values())}
    entries.sort(key=lambda item: (
        item["character"], proficiency_order[item["proficiency"]],
        rank_order.get(item["rank"], 99), item["skill_en"],
    ))
    unmapped = sorted({item["skill_en"] for item in entries if not item["skill_zh"]})
    payload = {
        "metadata": {
            "source_url": SOURCE_URL,
            "export_url": EXPORT_URL,
            "retrieved_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "character_count": len({item["character"] for item in entries}),
            "proficiency_count": len(PROFICIENCIES),
            "entry_count": len(entries),
            "unmapped_translation_count": len(unmapped),
            "sheet_character_counts": sheet_counts,
            "notes_zh": [
                "只擷取各分頁的角色列及熟練度等級欄；玩家留言、效果表及自動統計區沒有當成角色習得資料。",
                "中文名稱為本站暫譯，並保留英文原名供核對。",
                "原表帶問號的項目標為待確認；空白不等同角色不能習得。",
            ],
        },
        "proficiencies": list(PROFICIENCIES.values()),
        "unmapped_translations": unmapped,
        "entries": entries,
    }
    (ROOT / "character-learnsets.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        f"Wrote {len(entries)} entries for {payload['metadata']['character_count']} characters; "
        f"unmapped translations: {len(unmapped)}"
    )
    if unmapped:
        print("Unmapped:", " | ".join(unmapped))


if __name__ == "__main__":
    main()
